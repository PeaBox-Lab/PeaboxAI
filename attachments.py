from __future__ import annotations
import base64,io
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
TEXT_EXTENSIONS={'.txt','.md','.csv','.json','.xml','.html','.htm','.py','.ps1','.js','.ts','.yaml','.yml','.log','.ini','.cfg'}
IMAGE_EXTENSIONS={'.png','.jpg','.jpeg','.webp'}; MAX_EXTRACTED_CHARS=100000
def extension(name):return Path(name).suffix.lower()
def extract_uploaded_file(u):
 s=extension(u.name);raw=u.getvalue()
 if s in TEXT_EXTENSIONS:return raw.decode('utf-8',errors='replace')
 if s=='.pdf':return '\n\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(raw)).pages)
 if s=='.docx':
  d=Document(io.BytesIO(raw));out=[p.text for p in d.paragraphs]
  for table in d.tables:
   out += ['\t'.join(cell.text for cell in row.cells) for row in table.rows]
  return '\n'.join(out)
 if s in {'.xlsx','.xlsm'}:
  wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);out=[]
  for ws in wb.worksheets:
   out.append(f'### Worksheet: {ws.title}');out+=['\t'.join('' if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)]
  wb.close();return '\n'.join(out)
 if s=='.pptx':
  prs=Presentation(io.BytesIO(raw));return '\n'.join([f'### Slide {i}'+''.join('\n'+sh.text for sh in sl.shapes if hasattr(sh,'text') and sh.text) for i,sl in enumerate(prs.slides,1)])
 raise ValueError(f'Unsupported attachment type: {s or "unknown"}')
def build_user_content(prompt,uploads,supports_vision):
 sections=[];images=[];used=0
 for u in uploads or []:
  s=extension(u.name)
  if s in IMAGE_EXTENSIONS:
   if not supports_vision:raise ValueError(f'Selected model is not detected as vision-capable: {u.name}')
   media='image/jpeg' if s in {'.jpg','.jpeg'} else f'image/{s[1:]}';images.append({'name':u.name,'media_type':media,'base64':base64.b64encode(u.getvalue()).decode('ascii')});continue
  text=extract_uploaded_file(u);available=MAX_EXTRACTED_CHARS-used
  if available<=0:break
  text=text[:available];used+=len(text);sections.append(f'\n\n===== ATTACHMENT: {u.name} =====\n{text}')
 return prompt+('\n\nUse these attachments as source material:'+''.join(sections) if sections else ''),images
